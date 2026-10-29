"""Table OCR tool - local server.

Renders PDF/image pages, crops the user-selected region and sends it to a
vision model running in the local Ollama (WSL), then exports the edited table
to .xlsx. Everything runs on this machine; nothing leaves localhost.
"""
import base64
import http.client
import io
import json
import os
import re
import shutil
import socket
import subprocess
import sys
import threading
import time
import urllib.request
import uuid
import webbrowser
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from urllib.parse import parse_qs, unquote, urlparse

import pymupdf

import award
import award_batch
import mutation
import mutation_batch
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# --lan listens on all network interfaces so other devices on the local network can connect.
HOST = os.environ.get("TABLE_OCR_HOST", "0.0.0.0" if "--lan" in sys.argv else "127.0.0.1")
PORT = int(os.environ.get("TABLE_OCR_PORT", "8765"))
OLLAMA = os.environ.get("OLLAMA_URL", "http://localhost:11434")
DEFAULT_MODEL = os.environ.get("TABLE_OCR_MODEL", "qwen2.5vl:7b")
TRANSLATE_MODEL = os.environ.get("TABLE_OCR_TRANSLATE_MODEL", "translategemma:12b")
WSL_DISTRO = os.environ.get("TABLE_OCR_WSL_DISTRO", "")  # "" = default distro
WSL = ["wsl.exe"] + (["-d", WSL_DISTRO] if WSL_DISTRO else [])
STATIC = os.path.join(os.path.dirname(os.path.abspath(__file__)), "static")

DOCS = {}  # doc id -> pymupdf.Document
DOCS_LOCK = threading.Lock()

PROMPT = """You are a precise OCR engine for tables.
Transcribe the table in this image into JSON: {"rows": [[cell, cell, ...], ...]}.
Rules:
- One inner array per visual row, top to bottom, including the header row(s).
- Every row must have the same number of cells; use "" for empty cells.
- If a cell spans several columns, put its text in the first cell and "" in the others.
- Copy text exactly as printed (numbers, signs, decimals, currency symbols, %). Do not translate, sum, or invent values.
- Multi-line text inside one cell is joined with a single space.
Return only the JSON."""

SCHEMA = {
    "type": "object",
    "properties": {"rows": {"type": "array", "items": {"type": "array", "items": {"type": "string"}}}},
    "required": ["rows"],
}


# ---------------------------------------------------------------- Ollama ---

def ollama_request(path, payload=None, timeout=10):
    data = json.dumps(payload).encode() if payload is not None else None
    req = urllib.request.Request(OLLAMA + path, data=data, headers={"Content-Type": "application/json"})
    with urllib.request.urlopen(req, timeout=timeout) as r:
        return json.loads(r.read())


def ollama_up():
    try:
        ollama_request("/api/version", timeout=2)
        return True
    except Exception:
        return False


_KEEPALIVE = None
_ENSURE_LOCK = threading.Lock()


def keep_wsl_alive():
    """WSL shuts the Linux VM down when no Windows process is attached to it, taking Ollama with it
    mid-read. Hold one attached process open for as long as the tool runs, restarting it if it ended."""
    global _KEEPALIVE
    if sys.platform != "win32" or (_KEEPALIVE and _KEEPALIVE.poll() is None):
        return
    try:
        _KEEPALIVE = subprocess.Popen(WSL + ["-e", "sleep", "infinity"], creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0),
                                      stdin=subprocess.DEVNULL, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    except OSError:
        pass


def ensure_ollama():
    """Make sure the WSL Ollama is reachable; boot WSL / start `ollama serve` if needed."""
    with _ENSURE_LOCK:
        return _ensure_ollama()


def _wait_up(seconds):
    """Poll Ollama until it answers or `seconds` have passed (time-based: a probe itself can take ~2 s)."""
    end = time.time() + seconds
    while time.time() < end:
        if ollama_up():
            return True
        time.sleep(0.5)
    return False


def _ensure_ollama():
    wsl = WSL
    flags = getattr(subprocess, "CREATE_NO_WINDOW", 0)
    keep_wsl_alive()
    if _wait_up(15):  # WSL may be booting; Ollama starts with it
        return True
    if sys.platform == "win32":
        # Ollama is normally installed as a system service (its models live under the service's
        # account), so start that first; a plain `ollama serve` is only the last resort.
        try:
            subprocess.run(wsl + ["-u", "root", "-e", "systemctl", "start", "ollama"], timeout=60,
                           stdin=subprocess.DEVNULL, capture_output=True, creationflags=flags)
        except (OSError, subprocess.TimeoutExpired):
            pass
        if _wait_up(20):
            return True
        print("Ollama not reachable - starting `ollama serve` inside WSL ...")
        subprocess.Popen(wsl + ["-e", "bash", "-lc", "ollama serve"], creationflags=flags,
                         stdin=subprocess.DEVNULL, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
        return _wait_up(20)
    return False


# ------------------------------------------------------------ model health --
# Every request to the models is independent (no context builds up between pages), but over a long run
# the Ollama runtime can degrade — the two models nearly fill a 16 GB card — and calls start failing or
# coming back empty. A tiny test image with known printed text shows whether a model still reads properly.

CANARY_TEXT = ("3718", "2-14-0", "48/16/1/1")
_CANARY = None


def canary_png():
    global _CANARY
    if _CANARY is None:
        doc = pymupdf.open()
        page = doc.new_page(width=130, height=120)
        for i, text in enumerate(CANARY_TEXT):
            page.insert_text((14, 30 + i * 34), text, fontsize=17, fontname="helv")
        _CANARY = page.get_pixmap(dpi=220, alpha=False).tobytes("png")
    return _CANARY


def model_health(models, slow_after=90):
    """Read the test image with each model. -> {"ok", "detail", "seconds"}."""
    t_all = time.time()
    for model in models:
        t0 = time.time()
        try:
            res = ollama_request("/api/chat", {
                "model": model, "stream": False, "keep_alive": "15m",
                "options": {"temperature": 0, "num_ctx": 4096, "num_predict": 60},
                "messages": [{"role": "user", "content": mutation.PROMPT_PLAIN,
                              "images": [base64.b64encode(canary_png()).decode()]}],
            }, timeout=180)
        except Exception as e:
            return {"ok": False, "detail": f"{model} did not answer ({type(e).__name__})", "seconds": round(time.time() - t_all, 1)}
        seen = re.sub(r"\s+", "", mutation._norm(res.get("message", {}).get("content", "")))
        missing = [t for t in CANARY_TEXT if t not in seen]
        took = time.time() - t0
        if missing:
            return {"ok": False, "detail": f"{model} misread the test image (missed {', '.join(missing)})",
                    "seconds": round(time.time() - t_all, 1)}
        if took > slow_after:
            return {"ok": False, "detail": f"{model} took {took:.0f}s on the test image", "seconds": round(time.time() - t_all, 1)}
    return {"ok": True, "detail": "models read the test image correctly", "seconds": round(time.time() - t_all, 1)}


def unload_models(models):
    """Drop the models from memory; the next request loads them fresh."""
    for model in models:
        try:
            ollama_request("/api/generate", {"model": model, "keep_alive": 0}, timeout=60)
        except Exception:
            pass


def restart_ollama():
    """Restart the Ollama service inside WSL (a clean runtime), then wait for it."""
    if sys.platform == "win32":
        try:
            subprocess.run(WSL + ["-u", "root", "-e", "systemctl", "restart", "ollama"], timeout=90,
                           stdin=subprocess.DEVNULL, capture_output=True, creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0))
        except (OSError, subprocess.TimeoutExpired):
            pass
    time.sleep(2)
    return ensure_ollama()


_HEAL_LOCK = threading.Lock()


def heal_models(models, force=None):
    """Check the models and step up the remedy until they read the test image again:
    nothing -> unload and reload -> restart Ollama. force="soft"/"hard" skips the first check.
    -> {"ok", "action", "detail", "steps"}."""
    with _HEAL_LOCK:
        steps = []
        if not force:
            h = model_health(models)
            steps.append(f"check: {h['detail']}")
            if h["ok"]:
                return {"ok": True, "action": "none", "detail": h["detail"], "steps": steps}
        if force != "hard":
            unload_models(models)
            h = model_health(models)
            steps.append(f"reloaded models: {h['detail']}")
            if h["ok"]:
                return {"ok": True, "action": "reloaded", "detail": h["detail"], "steps": steps}
        restart_ollama()
        h = model_health(models)
        steps.append(f"restarted Ollama: {h['detail']}")
        return {"ok": h["ok"], "action": "restarted", "detail": h["detail"], "steps": steps}


def parse_rows(text):
    """Model output -> list of rows (list of str). Accepts JSON or a markdown table."""
    text = text.strip()
    m = re.search(r"\{.*\}|\[.*\]", text, re.S)
    if m:
        try:
            obj = json.loads(m.group(0))
            rows = obj.get("rows") if isinstance(obj, dict) else obj
            if isinstance(rows, list):
                return [[("" if c is None else str(c)).strip() for c in (r if isinstance(r, list) else [r])]
                        for r in rows]
        except json.JSONDecodeError:
            pass
    rows = []
    for line in text.splitlines():
        line = line.strip()
        if not line.startswith("|"):
            continue
        cells = [c.strip() for c in line.strip("|").split("|")]
        if all(re.fullmatch(r":?-{2,}:?", c) for c in cells if c):
            continue
        rows.append(cells)
    return rows


def normalize(rows):
    rows = [r for r in rows if any(c.strip() for c in r)]
    width = max((len(r) for r in rows), default=0)
    rows = [r + [""] * (width - len(r)) for r in rows]
    # drop columns that are empty everywhere
    keep = [i for i in range(width) if any(r[i].strip() for r in rows)]
    return [[r[i] for i in keep] for r in rows]


# ------------------------------------------------------------- documents ---

def open_document(data, filename):
    ext = os.path.splitext(filename)[1].lower().lstrip(".") or None
    doc = pymupdf.open(stream=data, filetype=ext)
    if not doc.is_pdf:  # images (png/jpg/tiff/...) -> convert so every page renders the same way
        pdf = pymupdf.open("pdf", doc.convert_to_pdf())
        doc.close()
        doc = pdf
    # Bake any /Rotate into the content so every page starts at 0; our own rotation is applied on render.
    for page in doc:
        if page.rotation:
            page.remove_rotation()
    return doc


def norm_rot(rot):
    return int(rot or 0) % 360 // 90 * 90


def unrotate_frac(rect_frac, rot):
    """Selection in the displayed (rotated clockwise by `rot`) page -> fractions of the unrotated page."""
    def pt(u, v):
        return {0: (u, v), 90: (v, 1 - u), 180: (1 - u, 1 - v), 270: (1 - v, u)}[rot]
    u0, v0, u1, v1 = rect_frac
    (ax, ay), (bx, by) = pt(u0, v0), pt(u1, v1)
    return min(ax, bx), min(ay, by), max(ax, bx), max(ay, by)


def render_page(doc, page_no, dpi, rot=0):
    zoom = dpi / 72
    return doc[page_no].get_pixmap(matrix=pymupdf.Matrix(zoom, zoom).prerotate(norm_rot(rot)), alpha=False)


def render_crop(doc, page_no, rect_frac, rot=0):
    rot = norm_rot(rot)
    page = doc[page_no]
    r = page.rect
    x0, y0, x1, y1 = unrotate_frac(rect_frac, rot)
    clip = pymupdf.Rect(r.x0 + x0 * r.width, r.y0 + y0 * r.height,
                        r.x0 + x1 * r.width, r.y0 + y1 * r.height)
    # Aim for the longest side at ~1600px: sharp enough for small print, small enough for the model.
    longest = max(clip.width, clip.height)
    zoom = max(1.0, min(1600 / longest, 6.0))
    pix = page.get_pixmap(matrix=pymupdf.Matrix(zoom, zoom).prerotate(rot), clip=clip, alpha=False)
    return pix.tobytes("png")


# ------------------------------------------------------- orientation check --

def orientation_from_text(page):
    """Digital PDFs: the direction of the text lines says exactly how far the page is turned."""
    votes = {0: 0, 90: 0, 180: 0, 270: 0}
    for block in page.get_text("dict")["blocks"]:
        for line in block.get("lines", []):
            dx, dy = line["dir"]
            n = sum(len(s["text"].strip()) for s in line["spans"])
            if abs(dx) >= abs(dy):
                votes[0 if dx > 0 else 180] += n
            else:  # text running down the page needs a 270° (anticlockwise) turn, and vice versa
                votes[270 if dy > 0 else 90] += n
    total = sum(votes.values())
    if total < 40:
        return None
    best = max(votes, key=votes.get)
    return best if votes[best] >= 0.6 * total else None


def is_sideways(page):
    """Scans: lines of text make dark/light stripes. Stronger stripes across columns than rows means the
    text runs vertically, i.e. the page is turned 90° one way or the other. None = can't tell."""
    zoom = 600 / max(page.rect.width, page.rect.height)
    pix = page.get_pixmap(matrix=pymupdf.Matrix(zoom, zoom), colorspace=pymupdf.csGRAY, alpha=False)
    w, h, s = pix.width, pix.height, pix.samples
    rows, cols = [0] * h, [0] * w
    for y in range(h):
        line = s[y * pix.stride: y * pix.stride + w]
        for x, v in enumerate(line):
            if v < 128:
                rows[y] += 1
                cols[x] += 1
    if sum(rows) < 0.002 * w * h:
        return None  # nearly blank

    def stripes(profile):
        return sum((a - b) ** 2 for a, b in zip(profile, profile[1:])) / max(1, len(profile))

    r, c = stripes(rows) / max(1, w), stripes(cols) / max(1, h)  # normalise by line length
    if c > 1.5 * r:
        return True
    if r > 1.5 * c:
        return False
    return None


_TESSERACT = None
_TESSERACT_CHECKED = 0.0


def tesseract_cmd():
    """Tesseract on PATH, else (on Windows) inside WSL. [] if unavailable. Vision LLMs can't reliably tell
    upright from upside-down; Tesseract's orientation detector can, for Latin and Arabic scripts alike.
    Only success is remembered: a failed check (e.g. while WSL is still booting) is retried 20 s later."""
    global _TESSERACT, _TESSERACT_CHECKED
    if _TESSERACT or time.time() - _TESSERACT_CHECKED < 20:
        return _TESSERACT or []
    _TESSERACT_CHECKED = time.time()
    if shutil.which("tesseract"):
        _TESSERACT = ["tesseract"]
    elif sys.platform == "win32":
        keep_wsl_alive()
        try:
            r = subprocess.run(WSL + ["-e", "which", "tesseract"], stdin=subprocess.DEVNULL, capture_output=True, timeout=45,
                               creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0))
            if r.returncode == 0:
                _TESSERACT = WSL + ["-e", "tesseract"]
        except (OSError, subprocess.TimeoutExpired):
            pass
    return _TESSERACT or []


def _osd(cmd, png):
    """One Tesseract orientation check -> (rotation, confidence) or None."""
    try:
        r = subprocess.run(cmd + ["stdin", "stdout", "--psm", "0"], input=png, capture_output=True, timeout=45,
                           creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0))
    except (OSError, subprocess.TimeoutExpired):
        return None
    out = r.stdout.decode(errors="replace")
    rot = re.search(r"Rotate:\s*(\d+)", out)
    conf = re.search(r"Orientation confidence:\s*([\d.]+)", out)
    return (norm_rot(rot.group(1)), float(conf.group(1))) if rot and conf else None


def orientation_from_tesseract(pngs):
    """Try each rendering (sizes differ in what Tesseract can make out); keep the most confident answer."""
    cmd = tesseract_cmd()
    if not cmd:
        return None
    best = None
    for png in pngs:
        res = _osd(cmd, png)
        if res and (best is None or res[1] > best[1]):
            best = res
        if best and best[1] >= 2.0:
            break  # confident enough
    return best[0] if best and best[1] >= 0.5 else None


def detect_orientation(doc, page_no, register=True):
    """-> (clockwise turn that makes the page upright, how it was decided).
    `register`: the page may be a (landscape) mutation register form; False for other documents (awards,
    letters), where only the text decides and a page is never turned by its shape alone."""
    with DOCS_LOCK:  # MuPDF isn't thread-safe; Tesseract runs outside the lock
        page = doc[page_no]
        rot = orientation_from_text(page)
        if rot is not None:
            return rot, "text"
        rot = mutation.orientation_from_grid(page) if register else None  # the form's own grid says which way up
        if rot is not None:
            return rot, "grid"
        pngs = []
        for longest in (2000, 1300, 2800):
            zoom = longest / max(page.rect.width, page.rect.height)
            pngs.append(page.get_pixmap(matrix=pymupdf.Matrix(zoom, zoom), colorspace=pymupdf.csGRAY, alpha=False).tobytes("png"))
        portrait = page.rect.height > page.rect.width
    rot = orientation_from_tesseract(pngs)
    if not register:
        return (rot, "tesseract") if rot is not None else (0, "unknown")
    if rot is not None and (rot in (90, 270)) == portrait:
        # A register page is a landscape form: portrait scans are on their side, landscape ones are
        # upright or upside down. A Tesseract answer that contradicts the page shape is a misread.
        return rot, "tesseract"
    if portrait:
        return 90, "guess"
    with DOCS_LOCK:
        sideways = is_sideways(doc[page_no])
    if sideways:
        # No sure answer from Tesseract. The register is a landscape form, so a page standing on its
        # side is turned; which way is a guess the user can flip with the rotate buttons.
        return 90, "guess"
    return 0, "unknown"


# ------------------------------------------------------------- translation --

# code -> (name used in the prompt, label shown in the UI)
LANGUAGES = {
    "ur": ("Urdu", "Urdu"),
    "pa-Arab": ("Punjabi", "Punjabi (Shahmukhi)"),
    "pa-Guru": ("Punjabi", "Punjabi (Gurmukhi)"),
    "sd-Arab": ("Sindhi", "Sindhi"),
    "en": ("English", "English"),
}
# Letters used in Sindhi but not in Urdu / Shahmukhi Punjabi.
SINDHI_LETTERS = set("ٻٽٿڀڃڄڇڊڌڍڏڙڦڪڳڱ")
PUNJABI_WORDS = {"دا", "دی", "دے", "نوں", "وچ", "توں", "تے", "ایہہ", "اوہ", "سی", "ہیگا", "کیہ", "جی", "اسیں", "تسیں"}
URDU_WORDS = {"کا", "کی", "کے", "کو", "میں", "سے", "ہے", "ہیں", "تھا", "یہ", "وہ", "اور", "نے"}


def detect_language(text):
    if re.search(r"[਀-੿]", text):
        return "pa-Guru"
    if any(ch in SINDHI_LETTERS for ch in text):
        return "sd-Arab"
    if not re.search(r"[؀-ۿ]", text):
        return "en"
    # Shahmukhi Punjabi shares Urdu's letters; tell them apart by common grammar words.
    words = re.findall(r"[؀-ۿݐ-ݿ]+", text)
    punjabi = sum(w in PUNJABI_WORDS for w in words) + 3 * text.count("ݨ")
    urdu = sum(w in URDU_WORDS for w in words)
    return "pa-Arab" if punjabi >= 2 and punjabi > urdu else "ur"


def translate_prompt(src, tgt, text):
    s_name, t_name = LANGUAGES[src][0], LANGUAGES[tgt][0]
    # Prompt format TranslateGemma was trained on (two blank lines before the text).
    return (f"You are a professional {s_name} ({src}) to {t_name} ({tgt}) translator. "
            f"Your goal is to accurately convey the meaning and nuances of the original {s_name} text "
            f"while adhering to {t_name} grammar, vocabulary, and cultural sensitivities.\n"
            f"Produce only the {t_name} translation, without any additional explanations or commentary. "
            f"Please translate the following {s_name} text into {t_name}:\n\n\n{text}")


# Reading and translating an image in one step makes the model invent content, so selections are
# transcribed first (it reads Urdu/Sindhi script well) and the text is then translated normally.
def transcribe_prompt(src):
    s_name = LANGUAGES[src][1] if src in LANGUAGES else "Urdu, Punjabi (Shahmukhi or Gurmukhi) or Sindhi"
    return (f"Transcribe all text in this image exactly as written, in its original script and language "
            f"(it may be {s_name}). Do not translate. Keep line and paragraph breaks. "
            f"Output only the transcribed text.")


def clean_transcription(text):
    text = re.sub(r"</?(start|end)_of_image>", "", text)
    lines = text.strip().splitlines()
    # The model sometimes prefixes a bare language label ("Punjabi") on its own line.
    if lines and lines[0].strip().lower() in {"urdu", "punjabi", "sindhi", "shahmukhi", "gurmukhi"}:
        lines = lines[1:]
    return "\n".join(lines).strip()


def chunk_text(text, limit=1500):
    """Split on paragraph (then sentence) boundaries so each request stays well inside the model's comfort zone."""
    chunks, cur = [], ""
    for para in re.split(r"\n\s*\n", text.strip()):
        pieces = [para] if len(para) <= limit else re.split(r"(?<=[.!?۔؟])\s+", para)
        for piece in pieces:
            if cur and len(cur) + len(piece) + 2 > limit:
                chunks.append(cur)
                cur = ""
            cur = f"{cur}\n\n{piece}" if cur else piece
    if cur:
        chunks.append(cur)
    return chunks


def ollama_generate(model, prompt, image=None):
    msg = {"role": "user", "content": prompt}
    if image:
        msg["images"] = [base64.b64encode(image).decode()]
    res = ollama_request("/api/chat", {
        "model": model, "stream": False, "keep_alive": "15m",
        "options": {"temperature": 0, "num_ctx": 8192, "num_predict": 4096},
        "messages": [msg],
    }, timeout=900)
    return res.get("message", {}).get("content", "").strip()


# ------------------------------------------------------------ excel export --

NUM_RE = re.compile(r"^\(?-?[$€£]?\s?\d{1,3}(,\d{3})*(\.\d+)?\)?%?$|^\(?-?[$€£]?\s?\d+(\.\d+)?\)?%?$")


def to_cell_value(s):
    t = s.strip()
    if not t or not NUM_RE.match(t) or (len(t) > 1 and t[0] == "0" and t[1].isdigit()):
        return s, None
    neg = t.startswith("(") and t.endswith(")") or t.lstrip("(").startswith("-")
    pct = t.endswith("%")
    core = re.sub(r"[()$€£%,\s-]", "", t)
    try:
        v = float(core)
    except ValueError:
        return s, None
    if neg:
        v = -v
    decimals = len(core.split(".")[1]) if "." in core else 0
    frac = ("." + "0" * decimals) if decimals else ""
    if pct:
        return v / 100, "0" + frac + "%"
    sym = next((ch for ch in "$€£" if ch in t), "")
    base = "#,##0" if "," in t or sym else "0"
    if not decimals:
        v = int(v)
    fmt = sym + base + frac  # keep printed separators/decimals (e.g. 1.50, $1,200)
    return v, (None if fmt == "0" else fmt)


MUTATION_HEADERS = ["S.No.", "Award /mutation No", "District", "Tehsil", "Mouza", "Khata Khatuni",
                    "Katat", "Kanal", "Marla", "Sq ft OR Sarsai", "Page"]  # Page = source page, for revision
MARLA_PER_KANAL = 20
THIRD_PER_MARLA = {"sarsai": 9, "sqft": 272}


def build_mutation_xlsx(rows, unit, issues, plot="katat"):
    """Sheet1: the fixed-header table. Totals: land total per award with live formulas pointing at Sheet1,
    so later edits in Excel carry through. Review: open issues, if any."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    head_font, fill = Font(bold=True), PatternFill("solid", fgColor="D9E1F2")
    headers = list(MUTATION_HEADERS)
    if plot == "khasra":
        headers[6] = "Khasra No."
    ws.append(headers)
    for c in ws[1]:
        c.font, c.fill = head_font, fill
    for i, row in enumerate(rows, start=1):
        vals = [i] + [str(v).strip() for v in (list(row[1:11]) + [""] * 10)[:10]]
        for j in (1, 6, 7, 8, 9, 10):  # numeric columns
            v = vals[j]
            if re.fullmatch(r"\d+", v):
                vals[j] = int(v)
            elif re.fullmatch(r"\d+\.\d+", v):
                vals[j] = float(v)
        ws.append(vals)
    for col, w in zip("ABCDEFGHIJK", (7, 20, 12, 14, 16, 28, 10, 8, 8, 16, 7)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"

    per = THIRD_PER_MARLA.get(unit, 9)
    third = "Sarsai" if unit != "sqft" else "Sq ft"
    groups = {}
    for row in rows:
        key = tuple(str(v).strip() for v in (row[1], row[2], row[3], row[4]))
        g = groups.setdefault(key, {"rows": 0})
        g["rows"] += 1
    if groups:
        t = wb.create_sheet("Totals")
        t.append(["Conversion rules", ""])
        t.append(["Marla per Kanal", MARLA_PER_KANAL])
        t.append([f"{third} per Marla", per])
        t.append([])
        hdr = ["Award /mutation No", "District", "Tehsil", "Mouza", "Rows", "Entered Kanal", "Entered Marla",
               f"Entered {third}", f"Total in {third}", "Total Kanal", "Total Marla", f"Total {third}",
               f"Total (Kanal-Marla-{third})"]
        t.append(hdr)
        for c in t[5]:
            c.font, c.fill = head_font, fill
        for n, (key, g) in enumerate(groups.items()):
            r = 6 + n
            award, district, tehsil, mouza = key
            crit = f'Sheet1!$B:$B,$A{r},Sheet1!$C:$C,$B{r},Sheet1!$D:$D,$C{r},Sheet1!$E:$E,$D{r}'
            t.append([int(award) if award.isdigit() else award, district, tehsil, mouza,
                      f"=COUNTIFS({crit})", f"=SUMIFS(Sheet1!$H:$H,{crit})", f"=SUMIFS(Sheet1!$I:$I,{crit})",
                      f"=SUMIFS(Sheet1!$J:$J,{crit})",
                      f"=F{r}*$B$2*$B$3+G{r}*$B$3+H{r}",
                      f"=INT(ROUND(I{r},4)/($B$2*$B$3))",
                      f"=INT(MOD(ROUND(I{r},4),$B$2*$B$3)/$B$3)",
                      f"=ROUND(MOD(ROUND(I{r},4),$B$3),2)",
                      f'=J{r}&"-"&K{r}&"-"&L{r}'])
        for col, w in zip("ABCDEFGHIJKLM", (20, 12, 14, 16, 7, 13, 13, 15, 15, 11, 11, 12, 26)):
            t.column_dimensions[col].width = w
    if issues:
        rv = wb.create_sheet("Review")
        rv.append(["S.No.", "Item"])
        for c in rv[1]:
            c.font, c.fill = head_font, fill
        for s_no, text in issues:
            rv.append([s_no, text])
        rv.column_dimensions["A"].width, rv.column_dimensions["B"].width = 8, 100
    wb.calculation.fullCalcOnLoad = True
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_xlsx(title, rows, header, merges=()):
    wb = Workbook()
    ws = wb.active
    ws.title = (re.sub(r"[\\/*?:\[\]]", "", title)[:31] or "Table")
    width = max((len(r) for r in rows), default=1)
    thin = Side(style="thin", color="A0A0A0")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    start = 1
    if title.strip():
        ws.cell(row=1, column=1, value=title.strip()).font = Font(bold=True, size=14)
        if width > 1:
            ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=width)
        ws.cell(row=1, column=1).alignment = Alignment(horizontal="center")
        start = 3
    for i, row in enumerate(rows):
        for j in range(width):
            text = row[j] if j < len(row) else ""
            is_head = header and i == 0
            value, fmt = (text, None) if is_head else to_cell_value(text)
            c = ws.cell(row=start + i, column=j + 1, value=value)
            c.border = border
            if fmt:
                c.number_format = fmt
            if is_head:
                c.font = Font(bold=True)
                c.fill = PatternFill("solid", fgColor="D9E1F2")
                c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    wide = set()  # text of cells merged across columns shouldn't widen just their first column
    for m in merges:
        r, c, rs, cs = int(m["r"]), int(m["c"]), int(m.get("rs", 1)), int(m.get("cs", 1))
        if r < 0 or c < 0 or r + rs > len(rows) or c + cs > width or (rs == 1 and cs == 1):
            continue
        ws.merge_cells(start_row=start + r, start_column=c + 1, end_row=start + r + rs - 1, end_column=c + cs)
        top = ws.cell(row=start + r, column=c + 1)
        top.alignment = Alignment(horizontal=top.alignment.horizontal or "center", vertical="center", wrap_text=True)
        if cs > 1:
            wide.add((r, c))
    for j in range(width):
        longest = max((len(r[j]) if j < len(r) and (i, j) not in wide else 0 for i, r in enumerate(rows)), default=8)
        ws.column_dimensions[get_column_letter(j + 1)].width = min(max(10, longest + 2), 60)
    if header and rows:
        # Freeze below the header, including header cells merged down into the next row(s).
        head_rows = max([1] + [int(m.get("rs", 1)) for m in merges if int(m["r"]) == 0])
        ws.freeze_panes = f"A{start + min(head_rows, len(rows))}"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ------------------------------------------------------------------ HTTP ---

# ------------------------------------------------------------- folder batches --
# Award and mutation folder conversions run in a background thread, one batch at a time (they share the
# models); the web tabs poll a job's state, answer its questions and show what it is building.

class Job:
    def __init__(self):
        self.lock = threading.Lock()
        self.answer = threading.Event()
        self.answer_path = None
        self.halt = threading.Event()
        self.state = {"running": False, "log": [], "root": "", "i": 0, "n": 0, "item": "", "question": None,
                      "live": None, "result": None}

    def log(self, line):
        with self.lock:
            self.state["log"].append(line)
            del self.state["log"][:-500]

    def update(self, **kw):
        with self.lock:
            self.state.update(kw)

    def snapshot(self):
        with self.lock:
            s = dict(self.state, log=self.state["log"][-200:])
        if s["question"]:
            s["question"] = dict(s["question"], seconds_left=max(0, int(s["question"]["deadline"] - time.time())))
        return s

    def ask(self, question, timeout):
        """Put a question to the web tab; -> the answer, or None after `timeout` seconds."""
        self.answer.clear()
        self.answer_path = None
        self.update(question=dict(question, deadline=time.time() + timeout))
        answered = self.answer.wait(timeout)
        self.update(question=None)
        return self.answer_path if answered else None

    def reply(self, path):
        self.answer_path = path
        self.answer.set()

    def stop(self):
        self.halt.set()
        self.answer.set()  # an open question takes its default

    def start(self, root, work):
        """Run work(job) in the background. Only one batch (of any kind) at a time."""
        with _JOBS_LOCK:
            if any(j.state["running"] for j in JOBS.values()):
                raise RuntimeError("A folder is already being converted - wait for it to finish or stop it.")
            if not os.path.isdir(root):
                raise RuntimeError(f"Folder not found: {root}")
            self.halt.clear()
            self.update(running=True, log=[], root=root, i=0, n=0, item="", question=None, live=None, result=None)

        def run():
            try:
                keep_wsl_alive()
                res = work(self)
            except Exception as e:
                self.log(f"!! The batch stopped: {e}")
                res = {"error": str(e)}
            self.update(running=False, question=None, result=res)

        threading.Thread(target=run, daemon=True).start()


JOBS = {"award": Job(), "mutation": Job()}
_JOBS_LOCK = threading.Lock()


def award_job_start(root, model, redo):
    def pick(folder, files):
        latest = max(files, key=award_batch.award_key)
        ans = job.ask({"folder": award_batch.folder_label(folder, root),
                       "files": [{"path": f, "name": os.path.basename(f), "latest": f == latest}
                                 for f in sorted(files, key=award_batch.award_key, reverse=True)]},
                      award_batch.PICK_TIMEOUT)
        if ans == "*all*":
            return list(files), award_batch.ALL_HOW
        if ans in files:
            return [ans], "chosen by user"
        return [latest], f"auto - no answer in {award_batch.PICK_TIMEOUT // 60} minutes, latest award taken"

    job = JOBS["award"]
    job.start(root, lambda j: award_batch.run(
        sys.modules[__name__], root, model, pick, j.log, redo=redo, stop=j.halt.is_set,
        step=lambda i, n, folder: j.update(i=i, n=n, item=award_batch.folder_label(folder, root))))


def register_ask(model, prompt, png):
    """One register cell read by a vision model (mutation.read_page's `ask`), retried when the
    connection drops (usually WSL shutting down under Ollama)."""
    body = {"model": model, "stream": False, "keep_alive": "15m",
            "options": {"temperature": 0, "num_ctx": 4096, "num_predict": 200},
            "messages": [{"role": "user", "content": prompt, "images": [base64.b64encode(png).decode()]}]}
    for attempt in range(3):
        try:
            return ollama_request("/api/chat", body, timeout=600).get("message", {}).get("content", "")
        except (OSError, http.client.HTTPException):
            if attempt == 2:
                raise
            time.sleep(2)
            if not ollama_up():
                ensure_ollama()


def mutation_job_start(root, model, redo):
    def ask_json(prompt, png, schema):  # the award's place names
        body = {"model": model, "stream": False, "format": schema, "keep_alive": "15m",
                "options": {"temperature": 0, "num_ctx": 8192, "num_predict": 3000},  # a whole page of entries
                "messages": [{"role": "user", "content": prompt, "images": [base64.b64encode(png).decode()]}]}
        return ollama_request("/api/chat", body, timeout=600).get("message", {}).get("content", "")

    def work(j):
        def live(doc, rows, page, pages):
            j.update(live={"doc": doc, "rows": rows, "page": page, "pages": pages, "headers": mutation_batch.HEADERS})

        return mutation_batch.run(sys.modules[__name__], root, (model, TRANSLATE_MODEL), register_ask, ask_json,
                                  j.log, redo=redo, stop=j.halt.is_set, live=live,
                                  step=lambda i, n, item: j.update(i=i, n=n, item=item))

    JOBS["mutation"].start(root, work)


def pick_folder_dialog():
    """Native folder picker on this PC (the browser can't hand out folder paths). -> path or ''"""
    import tkinter as tk
    from tkinter import filedialog
    win = tk.Tk()
    win.withdraw()
    win.attributes("-topmost", True)
    path = filedialog.askdirectory(parent=win, title="Folder to convert (subfolders are included)")
    win.destroy()
    return os.path.normpath(path) if path else ""


class Handler(BaseHTTPRequestHandler):
    def log_message(self, fmt, *args):
        sys.stderr.write("%s\n" % (fmt % args))

    def send(self, code, body, ctype="application/json", extra=None):
        if isinstance(body, (dict, list)):
            body = json.dumps(body).encode()
        elif isinstance(body, str):
            body = body.encode()
        self.send_response(code)
        self.send_header("Content-Type", ctype)
        self.send_header("Content-Length", str(len(body)))
        self.send_header("Cache-Control", "no-store")
        for k, v in (extra or {}).items():
            self.send_header(k, v)
        self.end_headers()
        self.wfile.write(body)

    def body(self):
        return self.rfile.read(int(self.headers.get("Content-Length", 0)))

    def get_doc(self, doc_id):
        with DOCS_LOCK:
            doc = DOCS.get(doc_id)
        if doc is None:
            raise KeyError("Document not found - please reopen the file.")
        return doc

    def do_GET(self):
        url = urlparse(self.path)
        q = {k: v[0] for k, v in parse_qs(url.query).items()}
        try:
            if url.path in ("/", "/index.html"):
                with open(os.path.join(STATIC, "index.html"), "rb") as f:
                    return self.send(200, f.read(), "text/html; charset=utf-8")
            if re.fullmatch(r"/[\w-]+\.(js|css)", url.path):  # flat static assets only (no paths)
                fpath = os.path.join(STATIC, url.path[1:])
                if os.path.isfile(fpath):
                    with open(fpath, "rb") as f:
                        ctype = "text/javascript" if fpath.endswith(".js") else "text/css"
                        return self.send(200, f.read(), ctype + "; charset=utf-8")
            if url.path == "/api/models":
                try:
                    tags = ollama_request("/api/tags", timeout=5)
                except Exception as e:
                    return self.send(200, {"ok": False, "error": f"Ollama not reachable at {OLLAMA}: {e}",
                                           "models": [], "default": DEFAULT_MODEL,
                                           "translate_default": TRANSLATE_MODEL})
                names = [m["name"] for m in tags.get("models", [])]
                return self.send(200, {"ok": True, "models": names, "default": DEFAULT_MODEL,
                                       "translate_default": TRANSLATE_MODEL})
            if url.path == "/api/version":
                return self.send(200, {"changed": code_changed()})
            if url.path == "/api/page":
                doc = self.get_doc(q["id"])
                with DOCS_LOCK:
                    png = render_page(doc, int(q["n"]), float(q.get("dpi", 150)), q.get("rot")).tobytes("png")
                return self.send(200, png, "image/png")
            if url.path == "/api/orient":
                doc = self.get_doc(q["id"])
                rot, how = detect_orientation(doc, int(q["n"]), register=q.get("form") != "page")
                return self.send(200, {"rotation": rot, "method": how})
            self.send(404, {"error": "not found"})
        except Exception as e:
            self.send(400, {"error": str(e)})

    def do_POST(self):
        url = urlparse(self.path)
        try:
            if url.path == "/api/upload":
                name = unquote(self.headers.get("X-Filename", "file.pdf"))
                doc = open_document(self.body(), name)
                doc_id = uuid.uuid4().hex
                with DOCS_LOCK:
                    DOCS[doc_id] = doc
                return self.send(200, {"id": doc_id, "pages": doc.page_count, "name": name})

            if url.path == "/api/extract":
                req = json.loads(self.body())
                doc = self.get_doc(req["id"])
                with DOCS_LOCK:
                    png = render_crop(doc, int(req["page"]), req["rect"], req.get("rot"))
                model = req.get("model") or DEFAULT_MODEL
                t0 = time.time()
                res = ollama_request("/api/chat", {
                    "model": model,
                    "stream": False,
                    "format": SCHEMA,
                    "keep_alive": "15m",
                    "options": {"temperature": 0, "num_ctx": 16384, "num_predict": 8192},
                    "messages": [{"role": "user", "content": PROMPT,
                                  "images": [base64.b64encode(png).decode()]}],
                }, timeout=900)
                text = res.get("message", {}).get("content", "")
                rows = normalize(parse_rows(text))
                if not rows:
                    return self.send(200, {"ok": False, "error": "The model did not return a table. Try a tighter selection.",
                                           "raw": text})
                return self.send(200, {"ok": True, "rows": rows, "seconds": round(time.time() - t0, 1),
                                       "model": model})

            if url.path == "/api/translate":
                req = json.loads(self.body())
                model = req.get("model") or TRANSLATE_MODEL
                tgt = req.get("target") or "en"
                src = req.get("source") or "auto"
                t0 = time.time()
                if req.get("rect"):  # selected region of the page: transcribe first, then translate
                    doc = self.get_doc(req["id"])
                    with DOCS_LOCK:
                        png = render_crop(doc, int(req["page"]), req["rect"], req.get("rot"))
                    text = clean_transcription(ollama_generate(model, transcribe_prompt(src), png))
                    if not text:
                        return self.send(200, {"ok": False, "error": "No text found in the selection."})
                else:
                    text = (req.get("text") or "").strip()
                    if not text:
                        return self.send(200, {"ok": False, "error": "Nothing to translate."})
                if src == "auto":
                    src = detect_language(text)
                if src == tgt:
                    return self.send(200, {"ok": False, "original": text,
                                           "error": f"Text looks like {LANGUAGES[src][1]} already - pick the source language."})
                out = "\n\n".join(ollama_generate(model, translate_prompt(src, tgt, c)) for c in chunk_text(text))
                return self.send(200, {"ok": True, "translation": out, "original": text, "source": src,
                                       "seconds": round(time.time() - t0, 1), "model": model})

            if url.path == "/api/mutation/read":
                req = json.loads(self.body())
                doc = self.get_doc(req["id"])
                reader = req.get("model") or DEFAULT_MODEL
                keep_wsl_alive()
                t0 = time.time()
                out = mutation.read_page(doc[int(req["page"])], norm_rot(req.get("rot")), register_ask, reader,
                                         TRANSLATE_MODEL, lock=DOCS_LOCK,
                                         plot_field="khasra" if req.get("plot") == "khasra" else "katat")  # held while rendering, released for model calls
                out["seconds"] = round(time.time() - t0, 1)
                return self.send(200, out)

            if url.path == "/api/mutation/export":
                req = json.loads(self.body())
                data = build_mutation_xlsx(req.get("rows", []), req.get("unit", "sarsai"), req.get("issues", []),
                                           req.get("plot", "katat"))
                return self.send(200, data, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                                 {"Content-Disposition": 'attachment; filename="mutation-register.xlsx"'})

            if url.path in ("/api/award/read", "/api/award/place"):
                # One award page: the land table's cells via the table extractor (or the place names
                # above it). The table itself is built from all pages by /api/award/build.
                req = json.loads(self.body())
                doc = self.get_doc(req["id"])
                model = req.get("model") or DEFAULT_MODEL
                keep_wsl_alive()

                def chat(prompt, png, schema):
                    body = {"model": model, "stream": False, "format": schema, "keep_alive": "30m",
                            "options": {"temperature": 0, "num_ctx": 16384, "num_predict": 8192},
                            "messages": [{"role": "user", "content": prompt, "images": [base64.b64encode(png).decode()]}]}
                    for attempt in range(3):
                        try:
                            return ollama_request("/api/chat", body, timeout=900).get("message", {}).get("content", "")
                        except (OSError, http.client.HTTPException):
                            if attempt == 2:
                                raise
                            time.sleep(2)
                            if not ollama_up():
                                ensure_ollama()

                t0 = time.time()
                with DOCS_LOCK:
                    page = doc[int(req["page"])]
                rot = norm_rot(req.get("rot"))
                if url.path == "/api/award/place":
                    place, filled = award.read_place(page, chat, rot), []
                    if any(not place.get(k) for k in ("district", "tehsil", "mouza")):
                        # no such columns in the table: the award names them in its opening text
                        with DOCS_LOCK:
                            first = doc[0]
                        place, filled = award.fill_place(place, award.read_place_doc(first, chat, norm_rot(req.get("rot0"))))
                    return self.send(200, {"place": place, "from_text": filled})
                cells = award.read_table(page, lambda png: parse_rows(chat(PROMPT, png, SCHEMA)), rot)
                return self.send(200, {"found": cells is not None, "cells": cells or [],
                                       "seconds": round(time.time() - t0, 1)})

            if url.path == "/api/award/build":
                req = json.loads(self.body())
                res = award.build([(int(n), cells) for n, cells in req.get("pages", [])])
                doc = DOCS.get(req.get("id"))
                cmd = tesseract_cmd() if doc is not None else None
                with DOCS_LOCK:
                    res["award"] = award.award_number(req.get("name", ""), doc[0] if doc is not None else None, cmd)
                return self.send(200, res)

            if url.path == "/api/award/browse":
                return self.send(200, {"path": pick_folder_dialog()})

            if url.path == "/api/award/scan":  # what a folder conversion would do, before starting it
                root = json.loads(self.body()).get("root", "").strip().strip('"')
                if not os.path.isdir(root):
                    return self.send(200, {"ok": False, "error": f"Folder not found: {root}"})
                folders = award_batch.find_folders(root)
                return self.send(200, {"ok": True, "folders": len(folders), "done": sum(1 for f in folders if f[2]),
                                       "multi": sum(1 for f in folders if len(f[1]) > 1)})

            if url.path == "/api/mutation/scan":  # what a mutation folder conversion would do
                root = json.loads(self.body()).get("root", "").strip().strip('"')
                if not os.path.isdir(root):
                    return self.send(200, {"ok": False, "error": f"Folder not found: {root}"})
                groups = mutation_batch.find_documents(root)
                docs = sum(len(f) for _d, f, _u in groups)
                done = sum(mutation_batch.is_done(p, os.listdir(d)) for d, f, u in groups for p in f + u)
                other = sum(len(u) for _d, _f, u in groups)  # named neither mutation nor award: checked when run
                return self.send(200, {"ok": True, "docs": docs, "other": other, "folders": len(groups), "done": done})

            m = re.fullmatch(r"/api/(award|mutation)/batch/(start|status|answer|stop)", url.path)
            if m:
                job, action = JOBS[m.group(1)], m.group(2)
                req = json.loads(self.body() or b"{}")
                if action == "start":
                    start = award_job_start if m.group(1) == "award" else mutation_job_start
                    start(req.get("root", "").strip().strip('"'), req.get("model") or DEFAULT_MODEL, bool(req.get("redo")))
                elif action == "status":
                    return self.send(200, job.snapshot())
                elif action == "answer":
                    job.reply(req.get("path"))
                else:
                    job.stop()
                return self.send(200, {"ok": True})

            if url.path == "/api/award/export":
                req = json.loads(self.body())
                data = award.build_xlsx(req.get("rows", []), req.get("info", []), req.get("issues", []))
                fname = re.sub(r'[\\/:*?"<>|]+', "-", req.get("filename") or "award-land").strip() + ".xlsx"
                return self.send(200, data, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                                 {"Content-Disposition": f'attachment; filename="{fname}"'})

            if url.path == "/api/export":
                req = json.loads(self.body())
                title = req.get("title", "")
                data = build_xlsx(title, req.get("rows", []), bool(req.get("header", True)), req.get("merges") or [])
                fname = (re.sub(r'[\\/:*?"<>|]+', "_", title).strip() or "table") + ".xlsx"
                return self.send(200, data,
                                 "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                                 {"Content-Disposition": f'attachment; filename="{fname}"'})

            if url.path == "/api/models/heal":
                # {"models": [...], "force": null|"soft"|"hard"} -> check, and reload/restart if needed
                req = json.loads(self.body() or b"{}")
                models = [m for m in (req.get("models") or [DEFAULT_MODEL, TRANSLATE_MODEL]) if m]
                return self.send(200, heal_models(models, req.get("force")))

            if url.path == "/api/models/unload":
                req = json.loads(self.body() or b"{}")
                unload_models([m for m in (req.get("models") or [DEFAULT_MODEL, TRANSLATE_MODEL]) if m])
                return self.send(200, {"ok": True})

            if url.path == "/api/close":
                req = json.loads(self.body())
                with DOCS_LOCK:
                    doc = DOCS.pop(req.get("id"), None)
                if doc:
                    doc.close()
                return self.send(200, {"ok": True})
            self.send(404, {"error": "not found"})
        except urllib.error.HTTPError as e:
            detail = e.read().decode(errors="replace")
            self.send(502, {"error": f"Ollama error {e.code}: {detail}"})
        except Exception as e:
            self.send(400, {"error": f"{type(e).__name__}: {e}"})


def lan_ips():
    ips = set()
    try:
        for info in socket.getaddrinfo(socket.gethostname(), None, socket.AF_INET):
            ips.add(info[4][0])
    except OSError:
        pass
    return sorted(ip for ip in ips if not ip.startswith(("127.", "169.254.", "172.")))  # skip loopback/WSL/Docker


CODE_FILES = [os.path.join(os.path.dirname(os.path.abspath(__file__)), f)
              for f in ("server.py", "award.py", "award_batch.py", "mutation.py", "mutation_batch.py")]
CODE_STAMP = {f: os.path.getmtime(f) for f in CODE_FILES if os.path.exists(f)}


def code_changed():
    """Python files edited since this server started (the running tool still uses the old code)."""
    return [os.path.basename(f) for f, t in CODE_STAMP.items() if os.path.exists(f) and os.path.getmtime(f) != t]


def main():
    url = f"http://127.0.0.1:{PORT}/"
    try:  # Windows lets a second server share the port, and the old one keeps answering: don't start twice
        with socket.create_connection(("127.0.0.1", PORT), timeout=1):
            print(f"Table OCR is already running at {url} - close that window first (or use it).", flush=True)
            if "--no-browser" not in sys.argv:
                webbrowser.open(url)
            return
    except OSError:
        pass

    def check_ollama():  # in the background, so the tool opens straight away while WSL/Ollama boot
        print(f"Checking Ollama at {OLLAMA} ...", flush=True)
        print("Ollama OK" if ensure_ollama() else "WARNING: Ollama not reachable; extraction will fail until it is up.", flush=True)
    threading.Thread(target=check_ollama, daemon=True).start()
    srv = ThreadingHTTPServer((HOST, PORT), Handler)
    print(f"Table OCR running at {url}  (Ctrl+C to stop)")
    if HOST != "127.0.0.1":
        print("Open from other devices on your network at:")
        for ip in lan_ips():
            print(f"    http://{ip}:{PORT}/")
    if "--no-browser" not in sys.argv:
        threading.Timer(0.8, lambda: webbrowser.open(url)).start()
    try:
        srv.serve_forever()
    except KeyboardInterrupt:
        pass


if __name__ == "__main__":
    main()
