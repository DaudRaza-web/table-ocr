"""Batch award extraction: every folder under a root, one award per folder -> an Excel land table.

    python award_batch.py ["D:\\...\\M3 Data"] [--redo] [--pages N] [--file award.pdf] [--out folder]

Each folder's award is read with the table extractor (qwen2.5vl only, every other model unloaded so it
has the machine's memory to itself), the land table is built (square carried down and joined to the
khasra, total K-M from the document) and saved next to the award as "Award <no> <mouza> - Land.xlsx".
When a folder holds more than one award, a window asks which to use; after 5 minutes without an answer
the latest award (highest year, then award number) is taken. Correction slips are not awards.
Folders that already have the Excel file are skipped unless --redo is given.
"""
import argparse
import base64
import datetime
import os
import re
import sys
import time
import traceback

ROOT = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, ROOT)
_argv, sys.argv = sys.argv, sys.argv[:1]  # server.py reads sys.argv at import
import award  # noqa: E402
import server  # noqa: E402
sys.argv = _argv

from openpyxl import Workbook  # noqa: E402
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side  # noqa: E402

MODEL = server.DEFAULT_MODEL
DEFAULT_ROOT = r"D:\Danish Enterprises\Motorway\M3 Data"
PICK_TIMEOUT = 5 * 60
OPTIONS = {"temperature": 0, "num_ctx": 16384, "num_predict": 8192}


# ------------------------------------------------------------------- model --

def one_model():
    """Unload every model except the table reader, then load it and keep it loaded."""
    server.ensure_ollama()
    try:
        loaded = [m["name"] for m in server.ollama_request("/api/ps").get("models", [])]
    except Exception:
        loaded = []
    server.unload_models([m for m in loaded if m != MODEL])
    server.ollama_request("/api/generate", {"model": MODEL, "keep_alive": "60m"}, timeout=600)


def _chat(prompt, png, schema):
    payload = {"model": MODEL, "stream": False, "format": schema, "keep_alive": "60m", "options": OPTIONS,
               "messages": [{"role": "user", "content": prompt, "images": [base64.b64encode(png).decode()]}]}
    for attempt in range(3):
        try:
            return server.ollama_request("/api/chat", payload, timeout=900).get("message", {}).get("content", "")
        except Exception:
            if attempt == 2:
                raise
            server.ensure_ollama()
            time.sleep(5)


def ask_table(png):
    return server.parse_rows(_chat(server.PROMPT, png, server.SCHEMA))


# ------------------------------------------------------------------ awards --

def is_award(name):
    n = name.lower()
    return n.endswith(".pdf") and "award" in n and "correction" not in n


def award_key(path):
    """Latest first: year, then award number, then file date."""
    m = award.AWARD_NO.search(os.path.basename(path))
    year = int(m.group(2)) if m and m.group(2) else 0
    no = int(m.group(1)) if m else 0
    return (year, no, os.path.getmtime(path))


def pick_award(folder, files):
    """Window asking which award to use; the latest is taken after PICK_TIMEOUT. -> (path, how)"""
    latest = max(files, key=award_key)
    try:
        import tkinter as tk
    except ImportError:
        return latest, "auto (no window available)"
    choice = {"path": latest, "how": f"auto - no answer in {PICK_TIMEOUT // 60} minutes, latest award taken"}
    win = tk.Tk()
    win.title("Which award?")
    win.attributes("-topmost", True)
    tk.Label(win, text=f"This folder has {len(files)} awards:\n{folder}\n\nWhich one should be used?",
             justify="left", padx=14, pady=10).pack(anchor="w")
    var = tk.StringVar(value=latest)
    for f in sorted(files, key=award_key, reverse=True):
        tk.Radiobutton(win, text=os.path.basename(f) + ("   (latest)" if f == latest else ""), variable=var,
                       value=f, anchor="w", padx=14).pack(fill="x")
    timer = tk.Label(win, padx=14, pady=6, fg="#555")
    timer.pack(anchor="w")
    end = time.time() + PICK_TIMEOUT

    def ok():
        choice.update(path=var.get(), how="chosen by user")
        win.destroy()

    def tick():
        left = int(end - time.time())
        if left <= 0:
            win.destroy()
            return
        timer.config(text=f"The latest award is used automatically in {left // 60}:{left % 60:02d}")
        win.after(1000, tick)

    tk.Button(win, text="Use this award", command=ok, padx=10).pack(pady=(0, 12))
    win.protocol("WM_DELETE_WINDOW", win.destroy)  # closing the window = take the latest
    tick()
    win.mainloop()
    return choice["path"], choice["how"]


# ------------------------------------------------------------------- excel --

HEADERS = ["S.No.", "Award No", "District", "Tehsil", "Locality/Mouza", "Khasra No.", "K", "M", "Page"]


def write_xlsx(path, src, res, how):
    wb = Workbook()
    ws = wb.active
    ws.title = "Award Land"
    thin = Side(style="thin", color="999999")
    box = Border(left=thin, right=thin, top=thin, bottom=thin)
    head = PatternFill("solid", fgColor="DDE6F0")
    ws.append(HEADERS)
    for c in ws[1]:
        c.font, c.fill, c.border = Font(bold=True), head, box
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    p = res["place"]
    for i, r in enumerate(res["rows"], 1):
        ws.append([i, res["award"], p.get("district", ""), p.get("tehsil", ""), p.get("mouza", ""),
                   r["khasra"], None, None, r["page"]])
    k, m = res["total"] or (None, None)
    ws.append(["", "", "", "", "", "Total", k, m, ""])
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.border = box
            c.alignment = Alignment(horizontal="center")
    for c in ws[ws.max_row]:
        c.font = Font(bold=True)
    for col, wdt in zip("ABCDEFGHI", (7, 11, 16, 16, 20, 14, 7, 7, 7)):
        ws.column_dimensions[col].width = wdt
    ws.freeze_panes = "A2"

    rv = wb.create_sheet("Review")
    rv.append(["Item", "Detail"])
    for c in rv[1]:
        c.font, c.fill = Font(bold=True), head
    rv.append(["Source file", src])
    rv.append(["Award chosen", how])
    rv.append(["Pages with the land table", ", ".join(map(str, res["pages"])) or "none"])
    rv.append(["Khasras listed", len(res["rows"])])
    rv.append(["Total (from the document)", f"{k} K - {m} M" if res["total"] else "not found"])
    rv.append(["Extracted", datetime.datetime.now().strftime("%Y-%m-%d %H:%M")])
    rv.append([])
    rv.append(["Check", "Issue"])
    for c in rv[rv.max_row]:
        c.font = Font(bold=True)
    issues = res["issues"]
    for it in issues or ["Nothing looked off in the extraction"]:
        rv.append(["!" if issues else "OK", it])
    rv.column_dimensions["A"].width = 26
    rv.column_dimensions["B"].width = 110
    if issues:
        rv.sheet_properties.tabColor = "E0A030"  # draws the eye to the Review tab
    wb.save(path)


def out_name(res):
    no = (res["award"] or "").replace("/", "-")
    mouza = re.sub(r'[<>:"/\\|?*]', "-", res["place"].get("mouza", "")).strip()
    return " ".join(x for x in ("Award", no, mouza, "- Land.xlsx") if x)


# -------------------------------------------------------------------- main --

def log(msg, fh=None):
    line = f"{datetime.datetime.now():%H:%M:%S}  {msg}"
    print(line, flush=True)
    if fh:
        fh.write(line + "\n")
        fh.flush()


def process(src, how, out_dir, max_pages, fh):
    cmd = server.tesseract_cmd()
    t0 = time.time()
    res = award.extract(src, ask_table, _chat, cmd, max_pages=max_pages,
                        progress=lambda n, tot: log(f"    page {n + 1}/{tot}", fh))
    path = os.path.join(out_dir, out_name(res))
    write_xlsx(path, src, res, how)
    tot = f"{res['total'][0]}-{res['total'][1]}" if res["total"] else "no total"
    log(f"  -> {os.path.basename(path)}: {len(res['rows'])} khasras, total {tot}, "
        f"{len(res['issues'])} issue(s), {time.time() - t0:.0f}s", fh)
    return path


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("root", nargs="?", default=DEFAULT_ROOT)
    ap.add_argument("--redo", action="store_true", help="re-extract folders that already have the Excel file")
    ap.add_argument("--pages", type=int, help="read only the first N pages (testing)")
    ap.add_argument("--file", help="one award PDF instead of the whole tree (testing)")
    ap.add_argument("--out", help="save the Excel file here instead of next to the award")
    a = ap.parse_args()

    log("Loading the table model (other models unloaded) ...")
    one_model()
    if a.file:
        process(a.file, "single file", a.out or os.path.dirname(a.file), a.pages, None)
        return
    fh = open(os.path.join(a.root, "Award Batch Log.txt"), "a", encoding="utf-8")
    log(f"=== Award batch over {a.root}", fh)
    done = failed = skipped = 0
    for folder, _dirs, names in sorted(os.walk(a.root)):
        files = [os.path.join(folder, n) for n in names if is_award(n)]
        if not files:
            continue
        rel = os.path.relpath(folder, a.root)
        if not a.redo and any(n.startswith("Award") and n.endswith("- Land.xlsx") for n in names):
            log(f"{rel}: already done, skipped", fh)
            skipped += 1
            continue
        if len(files) > 1:
            log(f"{rel}: {len(files)} awards - asking which to use", fh)
            src, how = pick_award(folder, files)
        else:
            src, how = files[0], "only award in the folder"
        log(f"{rel}: {os.path.basename(src)} ({how})", fh)
        try:
            process(src, how, a.out or folder, a.pages, fh)
            done += 1
        except Exception as e:
            failed += 1
            log(f"  !! failed: {e}", fh)
            fh.write(traceback.format_exc() + "\n")
    log(f"=== Finished: {done} extracted, {skipped} skipped, {failed} failed", fh)
    fh.close()


if __name__ == "__main__":
    main()
