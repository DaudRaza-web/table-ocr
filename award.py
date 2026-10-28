"""Award document extraction.

Award documents (Land Acquisition Act awards) are typed English. Their "detail of land" table has
District | Tehsil | Locality/Mouza (or Chack) | Sqr. No | Khasra (Killa) No. | K | M and can run over
several pages; the total is printed at the end ("264K - 10M" or a "Total:- 542 | 02" row).

Pipeline: find the table's ruled grid in code, crop its right-hand columns on each page and read them
with the table extractor (the same prompt as the Table tab), then build the land list in code:
the square number is carried down over blank rows and joined to the khasra (square 2, khasra 19/1 ->
"2/19/1"), plots printed in the Sqr column without a square ("116/1") are kept as they are, and the
stated total closes the table. Individual K/M areas are not reported; only the document's total is.
"""
import csv
import io
import json
import os
import re
import subprocess

import pymupdf

DPI = 300


# ------------------------------------------------------------------ helpers --

def _gray(page, rot=0, dpi=DPI):
    z = dpi / 72
    return page.get_pixmap(matrix=pymupdf.Matrix(z, z).prerotate(rot), colorspace=pymupdf.csGRAY, alpha=False)


def _dark(pix, thr=150):
    w, h, s, st = pix.width, pix.height, pix.samples, pix.stride
    return w, h, [bytes(1 if v < thr else 0 for v in s[y * st: y * st + w]) for y in range(h)]


def _group(idxs, gap=4):
    out = []
    for i in idxs:
        if out and i - out[-1][-1] <= gap:
            out[-1].append(i)
        else:
            out.append([i])
    return [(g[0], g[-1]) for g in out]


def tesseract(cmd, png, psm=6):
    """-> list of words {text, left, top, width, height, conf}."""
    r = subprocess.run(cmd + ["stdin", "stdout", "-l", "eng", "--psm", str(psm), "tsv"], input=png,
                       capture_output=True, timeout=120, creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0))
    rows = csv.DictReader(io.StringIO(r.stdout.decode(errors="replace")), delimiter="\t", quoting=csv.QUOTE_NONE)
    out = []
    for w in rows:
        if w.get("level") == "5" and (w.get("text") or "").strip():
            out.append({"text": w["text"].strip(), "left": int(w["left"]), "top": int(w["top"]),
                        "width": int(w["width"]), "height": int(w["height"]), "conf": float(w["conf"])})
    return out


# --------------------------------------------------------------------- grid --

def find_table(pix):
    """Ruled grid of the land table: vertical rules that run a long way down the page.
    -> {"v": [x...], "cols": last five rules, "y0", "y1", "size"} in pixels, or None."""
    w, h, D = _dark(pix, thr=215)  # table rules in some scans are light grey
    # longest vertical dark run per column (1 px of horizontal wobble allowed)
    best = [0] * w
    edge = int(0.015 * w)  # the scanned paper's own edge is not a rule
    for x in range(edge, w - edge):
        run = top = 0
        for y in range(h):
            if D[y][x] or D[y][x - 1] or D[y][x + 1]:
                run += 1
                if run > top:
                    top = run
            else:
                run = 0
        best[x] = top
    cand = [x for x in range(w) if best[x] > 0.12 * h]
    vs = [(a + b) // 2 for a, b in _group(cand, gap=6)]
    if len(vs) < 5:
        return None
    # The last four columns are always Sqr | Khasra | K | M (the District/Tehsil/Mouza columns on the
    # left are not ruled all the way down in every layout).
    right = vs[-5:]
    starts, ends = [], []
    for x in right:
        # the rule's longest stretch, bridging small breaks in the scan (up to 2.5% of the page)
        segs, s0, last = [], None, None
        for y in range(h):
            if D[y][x] or D[y][x - 1] or D[y][x + 1]:
                if s0 is None or y - last > 0.025 * h:
                    if s0 is not None:
                        segs.append((s0, last))
                    s0 = y
                last = y
        if s0 is not None:
            segs.append((s0, last))
        a, b = max(segs, key=lambda s: s[1] - s[0]) if segs else (0, 0)
        starts.append(a)
        ends.append(b)
    starts.sort()
    ends.sort()
    return {"v": vs, "cols": right, "y0": starts[2], "y1": ends[2], "size": (w, h)}


# -------------------------------------------------------------- read a page --

def read_table(page, ask, rot=0):
    """Land table cells on one page via the table extractor.
    `ask(png) -> list of rows (list of str)`. -> rows, or None when the page has no ruled table."""
    g = find_table(_gray(page, rot, dpi=200))
    if not g:
        return None
    w, h = g["size"]
    v, c = g["v"], g["cols"]
    # one column more on the left (Locality/Mouza) so a "Total" label printed there is read too
    left = v[-6] if len(v) >= 6 else c[0]
    r = page.rect
    rect = pymupdf.Rect((left / w - 0.005) * r.width, max(0, g["y0"] / h - 0.005) * r.height,
                        (c[4] / w + 0.005) * r.width, min(1, g["y1"] / h + 0.035) * r.height)
    z = 1800 / max(rect.width, rect.height)
    png = page.get_pixmap(matrix=pymupdf.Matrix(z, z).prerotate(rot), clip=rect, alpha=False).tobytes("png")
    return ask(png)


# ------------------------------------------------------------- build table --

NUM = re.compile(r"^\d{1,4}$")
KHASRA = re.compile(r"^\d{1,4}(/\d{1,4})*$")
TOTAL_KM = re.compile(r"(\d{1,5})\s*K\.?\s*[-–—~,]?\s*(\d{1,2})\s*M", re.I)
TOTAL_ROW = re.compile(r"total\D*(\d{1,5})\D+(\d{1,2})\b", re.I)


def _clean(t):
    t = (t or "").strip().strip(".,;:|[](){}_'\"`~-—").replace(" ", "")
    return t.replace("O", "0").replace("o", "0").replace("l", "1").replace("I", "1").replace("\\", "/")


def build(pages):
    """[(page_no, rows)] from the table extractor -> {"rows": [{"khasra", "page"}], "total", "issues"}.

    Per row: K and M are the last two plain numbers; before them come [square] and khasra. The table
    extractor sometimes leaves out blank cells, so a lone value counts as the khasra unless it sits in
    the Sqr column (a blank cell after it) or is above 30 with no square (killas run 1-25, sometimes
    26-27) - then it is a plot numbered without a square and kept as printed."""
    rows, issues, total, sqr = [], [], None, ""
    seen = {}
    for page_no, cells in pages:
        for raw in cells:
            joined = " ".join(x for x in raw if x)
            t = TOTAL_KM.search(joined) or TOTAL_ROW.search(joined)
            if t:
                total = (int(t.group(1)), int(t.group(2)))
                return {"rows": rows, "total": total, "issues": issues}
            # headings and place names are text; only numbers matter here
            # (a number misread with l/I/o/O in it is still a number)
            cl = ["" if re.search(r"[A-Za-z]", x) and not re.fullmatch(r"[\d/lIoO\s.]+", x) else _clean(x) for x in raw]
            idx = [i for i, x in enumerate(cl) if x]
            if not idx:
                continue
            nums = [i for i in idx if NUM.match(cl[i])]
            if len(nums) >= 2 and nums[-1] == idx[-1] and nums[-2] == idx[-2] and len(idx) >= 3:
                pre = cl[:nums[-2]]
                k, m = cl[nums[-2]], cl[nums[-1]]
            else:
                pre, k, m = cl, "", ""  # no K/M read on this line
                issues.append(f"Page {page_no}: no K/M area on line '{joined}' - check this row")
            vals = [(i, x) for i, x in enumerate(pre) if x]
            if not vals:
                continue
            plot = False
            if len(vals) >= 2:
                sq, kh = vals[0][1], vals[1][1]
                if len(vals) > 2:
                    issues.append(f"Page {page_no}: extra values on line '{joined}' - check this row")
            else:
                i, kh = vals[0]
                sq = ""
                in_sqr_col = len(pre) >= 2 and i == len(pre) - 2 and pre[-1] == ""
                first = re.match(r"\d+", kh)
                plot = in_sqr_col or (first is not None and int(first.group(0)) > 30)
            kh = kh.strip("/")
            if not KHASRA.match(kh) or (sq and not NUM.match(sq)):
                issues.append(f"Page {page_no}: unclear line '{joined}' - check the khasra number")
            if sq:
                sqr = sq.lstrip("0") or sq
            name = kh if plot or not sqr else f"{sqr}/{kh}"
            if not plot and not sqr and KHASRA.match(kh):
                issues.append(f"Page {page_no}: khasra {kh} has no square number above it")
            if name in seen:
                issues.append(f"Page {page_no}: khasra {name} appears twice (also on page {seen[name]}) - "
                              "a line may have been read twice")
            seen.setdefault(name, page_no)
            # k/m are kept only to take repeated khasras out of a combined total; they are not reported
            rows.append({"khasra": name, "page": page_no, "k": k, "m": m})
    return {"rows": rows, "total": total, "issues": issues}


# ------------------------------------------------------------ award header --

PLACE_SCHEMA = {"type": "object", "properties": {"district": {"type": "string"}, "tehsil": {"type": "string"},
                "locality": {"type": "string"}}, "required": ["district", "tehsil", "locality"]}
PLACE_PROMPT = ("This is the top-left of a typed land table. Its headings are District, Tehsil and Locality/Mouza "
                "(or Chack/Chak/Village). Copy the value printed under each heading exactly, as JSON "
                '{"district": "", "tehsil": "", "locality": ""}. Use "" for a heading that is not there.')


def read_place(page, ask_json, rot=0):
    """District / Tehsil / Locality from the table's top-left. `ask_json(prompt, png, schema) -> str`."""
    g = find_table(_gray(page, rot, dpi=200))
    if not g:
        return {}
    w, h = g["size"]
    r = page.rect
    y = g["y0"] / h
    rect = pymupdf.Rect(0.03 * r.width, max(0, y - 0.05) * r.height, (g["cols"][0] / w) * r.width, min(1, y + 0.07) * r.height)
    z = 1400 / max(rect.width, 1)
    png = page.get_pixmap(matrix=pymupdf.Matrix(z, z).prerotate(rot), clip=rect, alpha=False).tobytes("png")
    try:
        d = json.loads(ask_json(PLACE_PROMPT, png, PLACE_SCHEMA))
    except (ValueError, TypeError):
        return {}
    return {"district": d.get("district", "").strip(), "tehsil": d.get("tehsil", "").strip(),
            "mouza": d.get("locality", "").strip()}


DOC_PLACE_PROMPT = ("This is the first page of a land acquisition award. Find the district, the tehsil and the "
                    "village / mouza / chak (locality) that the award is for, as written on the page (e.g. "
                    '"Village Nazar Pakka, Tehsil & District Nankana Sahib"). Answer as JSON {"district": "", '
                    '"tehsil": "", "locality": ""} with just the names (no "District"/"Tehsil"/"Village" words); '
                    'use "" for one that is not on the page.')


def read_place_doc(page, ask_json, rot=0):
    """District / Tehsil / Locality from the text at the start of the award (for tables without those
    columns)."""
    r = page.rect
    rect = pymupdf.Rect(0, 0, r.width, r.height * 0.6)
    z = 1600 / max(rect.width, rect.height)
    png = page.get_pixmap(matrix=pymupdf.Matrix(z, z).prerotate(rot), clip=rect, alpha=False).tobytes("png")
    try:
        d = json.loads(ask_json(DOC_PLACE_PROMPT, png, PLACE_SCHEMA))
    except (ValueError, TypeError):
        return {}
    strip = lambda v, w: re.sub(rf"^\s*(?:{w})\.?\s*(?:no\.?)?\s*[:\-]?\s*", "", (v or "").strip(), flags=re.I).strip()  # noqa: E731
    return {"district": strip(d.get("district"), "district"), "tehsil": strip(d.get("tehsil"), "tehsil"),
            "mouza": strip(d.get("locality"), "village|mouza|mauza|locality")}


def fill_place(place, doc_place):
    """Place names from the table, with gaps filled from the document's opening text. -> (place, filled)"""
    out, filled = dict(place or {}), []
    for k in ("district", "tehsil", "mouza"):
        if not out.get(k) and (doc_place or {}).get(k):
            out[k] = doc_place[k]
            filled.append(k)
    return out, filled


AWARD_NO = re.compile(r"Award\s*No\.?\s*:?\s*(\d{1,3})\s*(?:[/-]\s*(\d{4}))?", re.I)


def award_number(path, page=None, cmd=None):
    """"12/2015" from the file name ("Award No. 12-2015 ..."), else from the page header by Tesseract."""
    m = AWARD_NO.search(os.path.basename(path))
    if m and m.group(2):
        return f"{int(m.group(1))}/{m.group(2)}"
    if page is not None and cmd:
        z = DPI / 72
        r = page.rect
        png = page.get_pixmap(matrix=pymupdf.Matrix(z, z), clip=pymupdf.Rect(0, 0, r.width, r.height * 0.15),
                              colorspace=pymupdf.csGRAY, alpha=False).tobytes("png")
        t = AWARD_NO.search(" ".join(wd["text"] for wd in tesseract(cmd, png, psm=6)))
        if t and t.group(2):
            return f"{int(t.group(1))}/{t.group(2)}"
    return f"{int(m.group(1))}" if m else ""


# ---------------------------------------------------------------- document --

def extract(path, ask_table, ask_json, cmd=None, max_pages=None, progress=None):
    """Whole award PDF -> {"award", "place", "rows": [{"khasra", "page"}], "total", "issues", "pages"}.
    `ask_table(png) -> rows`; `ask_json(prompt, png, schema) -> str`; `cmd` is Tesseract (award number)."""
    doc = pymupdf.open(path)
    n_pages = min(doc.page_count, max_pages or doc.page_count)
    pages, used, place, issues = [], [], {}, []
    for n in range(n_pages):
        if progress:
            progress(n, n_pages)
        cells = read_table(doc[n], ask_table)
        res = build(pages + [(n + 1, cells or [])])
        if not res["rows"] and res["total"] is None:
            if used:
                issues.append(f"Page {n + 1}: no land table here although the total was not reached yet")
                break
            continue  # the land table has not started
        pages.append((n + 1, cells or []))
        used.append(n + 1)
        if not place:
            place = read_place(doc[n], ask_json)
        if res["total"] is not None:
            break  # the total closes the land table
    res = build(pages)
    issues += res["issues"]
    if not used:
        issues.append("No land table (Sqr. No / Khasra No. / K / M) was found in the document")
    elif res["total"] is None:
        issues.append("No total K-M found at the end of the land table" +
                      (f" (only the first {n_pages} pages were read)" if max_pages else ""))
    labels = {"district": "District", "tehsil": "Tehsil", "mouza": "Locality/Mouza"}
    if doc.page_count and any(not place.get(k) for k in labels):
        # the table has no such columns: the award names them in its opening text
        place, filled = fill_place(place, read_place_doc(doc[0], ask_json))
        if filled:
            issues.append(f"{', '.join(labels[k] for k in filled)} not in the table - taken from the start of the "
                          "document; check")
    for k, label in labels.items():
        if not place.get(k):
            issues.append(f"{label} could not be read from the table or the document - fill it in")
    award = award_number(path, doc[0] if doc.page_count else None, cmd)
    if not award:
        issues.append("Award number not found in the file name or the page header")
    return {"award": award, "place": place, "rows": res["rows"], "total": res["total"], "issues": issues,
            "pages": used}


# ------------------------------------------------------------------- excel --

HEADERS = ["S.No.", "Award No", "District", "Tehsil", "Locality/Mouza", "Khasra No.", "K", "M", "Page"]


def table_rows(award, place, rows, total):
    """Rows of the land table under HEADERS; the K and M columns stay blank except on the closing Total row."""
    out = [[i, award, place.get("district", ""), place.get("tehsil", ""), place.get("mouza", ""), r["khasra"], "", "",
            r["page"]] for i, r in enumerate(rows, 1)]
    k, m = total or ("", "")
    return out + [["", "", "", "", "", "Total", k, m, ""]]


def _marla(k, m):
    return int(k) * 20 + int(m)


def _km(marla):
    return marla // 20, marla % 20


def combine(results):
    """Several awards of one place, oldest first -> one land list.

    A khasra already listed by an earlier award is left out (later awards and correction re-issues repeat
    land). Each award's own total stays as printed; the combined total is their sum less the areas of the
    khasras left out (read from the table, though not reported per khasra).
    -> {"rows": [{"award", "place", "khasra", "page"}], "totals": [(award, (K, M) or None)], "less": (K, M),
        "total": (K, M) or None, "ignored": [...], "issues": [...]}"""
    rows, ignored, issues, earlier, whole = [], [], [], {}, set()
    for res in results:
        a = res["award"] or "?"
        issues += [f"Award {a}: {t}" for t in res["issues"]]
        kept = []
        for r in res["rows"]:
            if r["khasra"] in earlier:
                ignored.append(dict(r, award=a, first=earlier[r["khasra"]], src=id(res)))
            else:
                kept.append(dict(r, award=a, place=res["place"]))
        rows += kept
        for r in kept:
            earlier.setdefault(r["khasra"], a)
        n_out = len(res["rows"]) - len(kept)
        if res["rows"] and not kept:
            issues.append(f"Award {a}: all {len(res['rows'])} khasras are already in an earlier award (the same award "
                          "scanned twice?) - none added, and its whole total is left out")
            if res["total"]:
                whole.add(id(res))  # leave out its printed total, not the sum of its rows as read
        elif n_out:
            issues.append(f"Award {a}: {n_out} khasra(s) already in an earlier award, left out: "
                          + ", ".join(f"{r['khasra']} (in {r['first']})" for r in ignored if r["award"] == a))
            # the left-out areas come from the table rows: if those don't add up to the award's printed
            # total, a row area was misread and the subtraction may be off too
            areas = [r for r in res["rows"] if NUM.match(r.get("k") or "") and NUM.match(r.get("m") or "")]
            got = sum(_marla(r["k"], r["m"]) for r in areas)
            if res["total"] and (len(areas) < len(res["rows"]) or got != _marla(*res["total"])):
                issues.append(f"Award {a}: its row areas add up to {_km(got)[0]}-{_km(got)[1]}, not its total "
                              f"{res['total'][0]}-{res['total'][1]}, so the repeats left out may be off - check "
                              "the left-out khasras' areas in the award")
    less = sum(_marla(*res["total"]) for res in results if id(res) in whole)
    for r in ignored:
        if r["src"] in whole:
            continue  # left out through its award's printed total above
        if NUM.match(r.get("k") or "") and NUM.match(r.get("m") or ""):
            less += _marla(r["k"], r["m"])
        else:
            issues.append(f"Award {r['award']}: the area of repeated khasra {r['khasra']} could not be read, so the "
                          "combined total still includes it - check")
    totals = [(res["award"] or "?", res["total"]) for res in results]
    total = None
    if all(t for _a, t in totals):
        total = _km(sum(_marla(*t) for _a, t in totals) - less)
    else:
        issues.append("Combined total not worked out: award(s) " + ", ".join(a for a, t in totals if not t)
                      + " have no total")
    return {"rows": rows, "totals": totals, "less": _km(less), "total": total, "ignored": ignored, "issues": issues}


def combined_rows(comb):
    """Rows under HEADERS for combine(): the khasras, each award's own total, the repeated khasras left
    out, and the combined Total."""
    out = [[i, r["award"], r["place"].get("district", ""), r["place"].get("tehsil", ""), r["place"].get("mouza", ""),
            r["khasra"], "", "", r["page"]] for i, r in enumerate(comb["rows"], 1)]
    for a, t in comb["totals"]:
        out.append(["", a, "", "", "", "Award total", *(t or ("", "")), ""])
    if comb["ignored"]:
        out.append(["", "", "", "", "", f"Less: {len(comb['ignored'])} repeated khasra(s)", *comb["less"], ""])
    out.append(["", "", "", "", "", "Total", *(comb["total"] or ("", "")), ""])
    return out


def build_xlsx(rows, info, issues, headers=HEADERS, widths=(7, 11, 16, 16, 20, 14, 7, 7, 7), text_cols=(1, 5),
               title="Award Land"):
    """rows: table_rows() (possibly edited); info: [(item, detail)] for the Review sheet; issues: [text].
    text_cols: columns kept as text (numbers like khasras); a row whose khasra column (the last text
    column) holds words - Total, Sub-total - is bold. -> .xlsx bytes"""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    def num(v):
        s = str(v if v is not None else "").strip()
        return int(s) if re.fullmatch(r"\d{1,6}", s) else s

    wb = Workbook()
    ws = wb.active
    ws.title = title
    thin = Side(style="thin", color="999999")
    box = Border(left=thin, right=thin, top=thin, bottom=thin)
    head = PatternFill("solid", fgColor="DDE6F0")
    ws.append(list(headers))
    for c in ws[1]:
        c.font, c.fill, c.border = Font(bold=True), head, box
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    label = max(text_cols)
    for r in rows:
        r = list(r) + [""] * (len(headers) - len(r))
        # khasra numbers stay text ("2/19/1"); S.No., areas and Page are numbers
        ws.append([str(v if v is not None else "") if i in text_cols else num(v) for i, v in enumerate(r)])
        bold = bool(re.search(r"[A-Za-z]", str(r[label])))  # Total / Sub-total / Award total / Less rows
        for c in ws[ws.max_row]:
            c.border = box
            c.alignment = Alignment(horizontal="center")
            if bold:
                c.font = Font(bold=True)
    from openpyxl.utils import get_column_letter
    for i, wdt in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = wdt
    ws.freeze_panes = "A2"

    rv = wb.create_sheet("Review")
    rv.append(["Item", "Detail"])
    for c in rv[1]:
        c.font, c.fill = Font(bold=True), head
    for item in info:
        rv.append(list(item))
    rv.append([])
    rv.append(["Check", "Issue"])
    for c in rv[rv.max_row]:
        c.font = Font(bold=True)
    for it in issues or ["Nothing looked off in the extraction"]:
        rv.append(["!" if issues else "OK", it])
    rv.column_dimensions["A"].width = 26
    rv.column_dimensions["B"].width = 110
    if issues:
        rv.sheet_properties.tabColor = "E0A030"  # draws the eye to the Review tab
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
